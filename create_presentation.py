import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color scheme
NAVY = RGBColor(20, 45, 78)
BLUE = RGBColor(38, 89, 140)
LIGHT = RGBColor(245, 248, 252)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(40, 40, 40)
GREEN = RGBColor(40, 167, 69)
RED = RGBColor(220, 53, 69)

def add_title_slide(title, subtitle):
    """Add title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = NAVY
    background.line.fill.background()
    slide.shapes._spTree.remove(background._element)
    slide.shapes._spTree.insert(2, background._element)
    
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.5))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Calibri"
    
    subtitle_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(11.7), Inches(1))
    p2 = subtitle_box.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(28)
    p2.font.color.rgb = LIGHT
    p2.font.name = "Calibri"

def add_content_slide(title, content_list):
    """Add content slide with bullets"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = WHITE
    background.line.fill.background()
    slide.shapes._spTree.remove(background._element)
    slide.shapes._spTree.insert(2, background._element)
    
    # Header
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1))
    header.fill.solid()
    header.fill.fore_color.rgb = BLUE
    header.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.2), Inches(12), Inches(0.7))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Calibri"
    
    # Content
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, item in enumerate(content_list):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        p.text = item
        p.font.size = Pt(24)
        p.font.name = "Calibri"
        p.font.color.rgb = DARK
        p.space_after = Pt(14)
        p.level = 0

# Slide 1: Title
add_title_slide(
    "AI-Powered Defect Triage & RCA Tool",
    "Real-time Engineering Defect Investigation"
)

# Slide 2: What It Does
add_content_slide(
    "What This Tool Does",
    [
        "✅ Accepts Jira defect input",
        "✅ Predicts owner team (80-95% accuracy)",
        "✅ Identifies affected subsystem",
        "✅ Suggests probable root causes",
        "✅ Finds similar historical defects",
        "✅ Generates investigation report",
    ]
)

# Slide 3: How It Works
add_content_slide(
    "Live Demo: Cluster Wakeup Defect",
    [
        "Input: 'Cluster wakeup failure after ignition on'",
        "",
        "Output:",
        "  • Owner: Cluster Team (80%)",
        "  • Subsystem: Instrument Cluster",
        "  • Root Cause: Missing NM Signal (94% confidence)",
        "  • Similar Defects: 18 found",
        "  • Est. Investigation Time: 1.5-3 hours",
    ]
)

# Slide 4: Business Value
add_content_slide(
    "Business Value",
    [
        "⏱️  Reduces triage time from hours to minutes",
        "",
        "👥 Reduces re-routing across teams",
        "",
        "📊 Improves consistency in early investigation",
        "",
        "📚 Leverages historical defect knowledge",
        "",
        "🎯 Gets engineers to debugging faster",
    ]
)

# Slide 5: Call to Action
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
background.fill.solid()
background.fill.fore_color.rgb = LIGHT
background.line.fill.background()
slide.shapes._spTree.remove(background._element)
slide.shapes._spTree.insert(2, background._element)

title_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.3), Inches(1))
p = title_box.text_frame.paragraphs[0]
p.text = "Try It Now"
p.font.size = Pt(48)
p.font.bold = True
p.font.color.rgb = NAVY
p.font.name = "Calibri"

info_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.8), Inches(10.3), Inches(3.5))
tf = info_box.text_frame
tf.word_wrap = True

items = [
    "🔗 GitHub: github.com/1998NiKi/Defect-Triage-RCA-Stand-alone",
    "",
    "🚀 Run locally: streamlit run app.py",
    "",
    "📦 Try sample inputs for instant results",
]

for i, item in enumerate(items):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.text = item
    p.font.size = Pt(22)
    p.font.name = "Calibri"
    p.font.color.rgb = DARK
    p.space_after = Pt(12)

# Save
output_path = "Defect_Triage_RCA_Presentation_v2.pptx"
prs.save(output_path)
print(f"✅ Presentation created: {output_path}")
